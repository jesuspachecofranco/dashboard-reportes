import calendar
from datetime import datetime
import io
import os
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import streamlit as st

# Configuración de la página con ancho completo
st.set_page_config(
    page_title="Dashboard de Reportes 2026", page_icon="📊", layout="wide"
)

st.title("📊 Control y Seguimiento de Incidencias")

# Inicializar variables de estado si no existen
if "lote_seleccionado" not in st.session_state:
    st.session_state["lote_seleccionado"] = pd.DataFrame()

if "lotes_guardados" not in st.session_state:
    st.session_state["lotes_guardados"] = {}  # Diccionario: {nombre_lote: DataFrame_congelado}


# ==========================================
# FUNCIÓN DE DEPURACIÓN Y UNIFICACIÓN DE TXT (CON PROTECCIÓN)
# ==========================================
def procesar_txts_seguro(archivos_subidos):
    mapeo_zonas = {
        "01": "NORTE",
        "02": "SUR",
        "03": "ESTE",
        "04": "OESTE",
        "05": "CENTRO",
        "06": "VÍA DUACA",
        "07": "VÍA RÍO CLARO",
        "08": "VÍA PAVIA",
        "09": "VÍA BUENA VISTA",
        "10": "VIA VIEJA CARORA",
        "11": "VÍA QUIBOR",
        "12": "VÍA AUTOPISTA CARORA",
    }

    meses_es = {
        1: "ENERO",
        2: "FEBRERO",
        3: "MARZO",
        4: "ABRIL",
        5: "MAYO",
        6: "JUNIO",
        7: "JULIO",
        8: "AGOSTO",
        9: "SEPTIEMBRE",
        10: "OCTUBRE",
        11: "NOVIEMBRE",
        12: "DICIEMBRE",
    }

    indices_a_conservar = [
        13,
        14,
        15,
        16,
        17,
        18,
        19,
        20,
        21,
        22,
        23,
        25,
        26,
        28,
        29,
        30,
        31,
        32,
        34,
        35,
        36,
        37,
        38,
    ]
    nombres_finales = [
        "INCIDENCIA",
        "DIRECCIÓN",
        "MOTIVO",
        "RECEPCIÓN",
        "FINALIZACIÓN",
        "TEC",
        "TCR",
        "TDV",
        "TEJ",
        "TAR",
        "ESTADO",
        "CLIENTE",
        "TOTAL INCIDENCIAS",
        "SUMATORIAS TEC",
        "SUMATORIAS TCR",
        "SUMATORIAS TDV",
        "SUMATORIAS TEJ",
        "SUMATORIAS TAR",
        "PROMEDIO TEC",
        "PROMEDIO TCR",
        "PROMEDIO TDV",
        "PROMEDIO TEJ",
        "PROMEDIO TAR",
    ]

    lista_dataframes = []

    for archivo in archivos_subidos:
        nombre_archivo = archivo.name
        try:
            contenido = archivo.getvalue().decode("latin-1")
            df = pd.read_csv(
                io.StringIO(contenido),
                sep="|",
                header=None,
                dtype=str,
                engine="python",
                on_bad_lines="skip",
            )

            total_columnas_archivo = df.shape[1]
            max_indice_requerido = max(indices_a_conservar)

            if total_columnas_archivo <= max_indice_requerido:
                continue

            df = df.iloc[:, indices_a_conservar]
            df.columns = nombres_finales

            codigo = nombre_archivo[:2]
            nombre_zona = mapeo_zonas.get(codigo, nombre_archivo)
            df.insert(0, "ZONA", nombre_zona)

            for col in df.columns:
                df[col] = df[col].astype(str).str.strip().str.upper()
                df[col] = df[col].replace(
                    ["NAN", "NONE", "NAT", "NAN.0"], "", regex=False
                )
                if col == "INCIDENCIA":
                    df[col] = df[col].str.replace(r"\.0$", "", regex=True)

            df = df.replace(r"\*", "", regex=True)
            df = df[df["INCIDENCIA"] != ""]
            df = df[df["INCIDENCIA"].str.match(r"^\d+$", na=False)]
            df = df[df["RECEPCIÓN"] != ""]

            if df.empty:
                continue

            dt_recepcion = pd.to_datetime(
                df["RECEPCIÓN"], format="%d-%m-%y %I:%M %p", errors="coerce"
            )
            df["_DT_RECEPCION_TEMP"] = dt_recepcion
            df = df.dropna(subset=["_DT_RECEPCION_TEMP"])
            df = df.drop(columns=["_DT_RECEPCION_TEMP"])

            dt_recepcion = pd.to_datetime(
                df["RECEPCIÓN"], format="%d-%m-%y %I:%M %p", errors="coerce"
            )
            mes_recibido = dt_recepcion.dt.month.map(meses_es).fillna("")
            anio_recibido = (
                dt_recepcion.dt.year.fillna("")
                .astype(str)
                .str.replace(r"\.0$", "", regex=True)
            )

            dt_finalizacion = pd.to_datetime(
                df["FINALIZACIÓN"], format="%d-%m-%y %I:%M %p", errors="coerce"
            )
            mes_finalizado = dt_finalizacion.dt.month.map(meses_es).fillna("")
            anio_finalizado = (
                dt_finalizacion.dt.year.fillna("")
                .astype(str)
                .str.replace(r"\.0$", "", regex=True)
            )

            idx_rec = df.columns.get_loc("RECEPCIÓN")
            df.insert(idx_rec, "MES RECIBIDO", mes_recibido)
            df.insert(idx_rec, "AÑO RECIBIDO", anio_recibido)

            idx_fin = df.columns.get_loc("FINALIZACIÓN")
            df.insert(idx_fin, "MES FINALIZADO", mes_finalizado)
            df.insert(idx_fin, "AÑO FINALIZADO", anio_finalizado)

            df = df.drop_duplicates()
            lista_dataframes.append(df)
        except Exception as e:
            st.error(
                f"⚠️ Error procesando el archivo '{nombre_archivo}': {e}"
            )
            return False

    if lista_dataframes:
        df_nuevo = pd.concat(lista_dataframes, ignore_index=True)
        archivo_actualizacion = "resultado_actualizacion.xlsx"
        df_nuevo.to_excel(archivo_actualizacion, index=False)

        if os.path.exists("resultado_unificado.xlsx"):
            try:
                timestamp_respaldo = datetime.now().strftime("%Y%m%d_%H%M%S")
                archivo_respaldo = (
                    f"resultado_unificado_respaldo_{timestamp_respaldo}.xlsx"
                )
                df_existente_seguridad = pd.read_excel(
                    "resultado_unificado.xlsx", dtype=str
                )
                df_existente_seguridad.to_excel(archivo_respaldo, index=False)

                df_final = pd.concat(
                    [df_existente_seguridad, df_nuevo], ignore_index=True
                )
                df_final = df_final.drop_duplicates(
                    subset=["INCIDENCIA"], keep="last"
                )
            except Exception:
                df_final = df_nuevo
        else:
            df_final = df_nuevo

        wb = Workbook()
        ws = wb.active
        ws.title = "Consolidado Zonas"
        ws.views.sheetView[0].showGridLines = True

        headers = list(df_final.columns)
        ws.append(headers)

        for row in df_final.itertuples(index=False, name=None):
            ws.append(list(row))

        max_row = ws.max_row
        max_col = len(headers)
        end_cell = get_column_letter(max_col) + str(max_row)

        tab = Table(displayName="TablaIncidencias", ref=f"A1:{end_cell}")
        style = TableStyleInfo(
            name="TableStyleLight1",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)
        wb.save("resultado_unificado.xlsx")
        return True
    return False


# ==========================================
# CARGA Y VISUALIZACIÓN DE DATOS
# ==========================================
@st.cache_data
def cargar_datos():
    archivo_entrada = "resultado_unificado.xlsx"
    if not os.path.exists(archivo_entrada):
        return None

    df = pd.read_excel(archivo_entrada, dtype=str)
    df["RECEPCION_DT"] = pd.to_datetime(
        df["RECEPCIÓN"], format="%d-%m-%y %I:%M %p", errors="coerce"
    )
    df["FINALIZACION_DT"] = pd.to_datetime(
        df["FINALIZACIÓN"], format="%d-%m-%y %I:%M %p", errors="coerce"
    )
    return df


df = cargar_datos()

if df is None:
    st.warning(
        "⚠️ No se encontró el archivo 'resultado_unificado.xlsx'. Ve a la"
        " pestaña **📁 Actualización** para subir tus archivos .txt"
        " iniciales."
    )

columnas_mostrar = [
    "INCIDENCIA",
    "RECEPCIÓN",
    "FINALIZACIÓN",
    "DIRECCIÓN",
    "CLIENTE",
    "MOTIVO",
    "ESTADO",
]

# ==========================================
# VISTAS PRINCIPALES EN TABS (6 TABS EN TOTAL)
# ==========================================
tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
    "📅 Seguimiento Diario",
    "📈 Seguimiento Anual",
    "🔍 Búsqueda Avanzada",
    "🎯 Análisis",
    "📄 Reportes",
    "📦 Actualización",
])


def renderizar_tabla_con_seleccion(df_sub, nombre_pestana, tipo_origen_etiqueta=None):
    if df is None:
        st.info("Primero debes cargar datos en la pestaña 'Carga y Actualización'.")
        return
    if df_sub.empty:
        st.info("No hay registros para mostrar en esta sección.")
        return

    df_editable = df_sub[columnas_mostrar].copy()
    df_editable.insert(0, "Seleccionar", False)

    df_editado = st.data_editor(
        df_editable,
        use_container_width=True,
        key=f"editor_{nombre_pestana}",
        hide_index=True,
    )

    col_btn1, col_btn2 = st.columns([2, 4])
    with col_btn1:
        if st.button(
            f"➕ Agregar al Lote de Análisis",
            key=f"btn_agregar_{nombre_pestana}",
        ):
            seleccionadas = df_editado[df_editado["Seleccionar"] == True]
            if not seleccionadas.empty:
                ids_seleccionados = seleccionadas["INCIDENCIA"].tolist()
                df_lote_nuevo = df[
                    df["INCIDENCIA"].isin(ids_seleccionados)
                ].copy()

                if tipo_origen_etiqueta:
                    df_lote_nuevo["ORIGEN_TIPO"] = tipo_origen_etiqueta
                else:
                    df_lote_nuevo["ORIGEN_TIPO"] = "Recibidas en el Día"

                if "UNIDAD" not in df_lote_nuevo.columns:
                    df_lote_nuevo["UNIDAD"] = ""
                if "GERENCIA" not in df_lote_nuevo.columns:
                    df_lote_nuevo["GERENCIA"] = ""

                if not st.session_state["lote_seleccionado"].empty:
                    df_acumulado = pd.concat([
                        st.session_state["lote_seleccionado"],
                        df_lote_nuevo,
                    ]).drop_duplicates(subset=["INCIDENCIA"])
                    st.session_state["lote_seleccionado"] = df_acumulado
                else:
                    st.session_state["lote_seleccionado"] = df_lote_nuevo

                st.success(
                    f"¡Se agregaron {len(ids_seleccionados)} incidencias a la pestaña Análisis! (Total lote: {len(st.session_state['lote_seleccionado'])})"
                )
            else:
                st.warning("⚠️ No has seleccionado ninguna incidencia en la tabla.")


# Función auxiliar para renderizar gráficos de un DataFrame de lote dado (con sufijo anti-duplicados)
def renderizar_panel_graficos(df_lote_objetivo, sufijo=""):
    col_g1, col_g2 = st.columns(2)

    with col_g1:
        st.markdown("##### 📍 1. Gráfico de Zonas")
        if "ZONA" in df_lote_objetivo.columns and not df_lote_objetivo["ZONA"].isna().all():
            fig_zona = px.pie(df_lote_objetivo, names="ZONA", title="Distribución por Zonas", hole=0.4)
            st.plotly_chart(fig_zona, use_container_width=True, key=f"fig_zona_{sufijo}")
        else:
            st.info("No hay datos de zona disponibles en este lote.")

    with col_g2:
        st.markdown("##### 📅 2. Gráfico por Mes de Recepción")
        if "RECEPCION_DT" in df_lote_objetivo.columns and not df_lote_objetivo["RECEPCION_DT"].isna().all():
            df_meses_lote = df_lote_objetivo.groupby(df_lote_objetivo["RECEPCION_DT"].dt.to_period("M")).size().reset_index(name="CANTIDAD")
            df_meses_lote["MES_STR"] = df_meses_lote["RECEPCION_DT"].dt.strftime("%B %Y").str.upper()

            fig_mes = px.bar(
                df_meses_lote,
                x="MES_STR",
                y="CANTIDAD",
                title="Incidencias Agrupadas por Mes de Recepción",
                labels={"MES_STR": "Mes", "CANTIDAD": "Cantidad"},
                color="CANTIDAD",
                color_continuous_scale="Blues",
            )
            st.plotly_chart(fig_mes, use_container_width=True, key=f"fig_mes_{sufijo}")
        else:
            st.info("No hay fechas válidas para este gráfico.")

    col_g3, col_g4 = st.columns(2)

    with col_g3:
        st.markdown("##### 🏷️ 3. Gráfico de Motivos")
        if "MOTIVO" in df_lote_objetivo.columns:
            fig_motivo = px.bar(
                df_lote_objetivo["MOTIVO"].value_counts().reset_index(),
                x="MOTIVO",
                y="count",
                title="Incidencias por Grupo de Origen / Motivo",
                labels={"MOTIVO": "Motivo", "count": "Cantidad"},
            )
            st.plotly_chart(fig_motivo, use_container_width=True, key=f"fig_motivo_{sufijo}")

    with col_g4:
        st.markdown("##### 🔄 4. Gráfico por Origen")
        if "ORIGEN_TIPO" in df_lote_objetivo.columns and not df_lote_objetivo["ORIGEN_TIPO"].isna().all():
            fig_origen = px.pie(
                df_lote_objetivo,
                names="ORIGEN_TIPO",
                title="Clasificación por Origen",
                hole=0.4,
                color="ORIGEN_TIPO",
                color_discrete_map={"Acumuladas": "#94a3b8", "Recibidas en el Día": "#3b82f6"}
            )
            st.plotly_chart(fig_origen, use_container_width=True, key=f"fig_origen_{sufijo}")
        else:
            st.info("No se cuenta con la clasificación de origen en este lote.")

    col_g5, _ = st.columns(2)
    with col_g5:
        st.markdown("##### 🏢 5. Gráfico por Unidad Asignada")
        df_unidad_valida = df_lote_objetivo[(df_lote_objetivo["UNIDAD"].notna()) & (df_lote_objetivo["UNIDAD"] != "")]
        if not df_unidad_valida.empty:
            fig_unidad = px.bar(
                df_unidad_valida["UNIDAD"].value_counts().reset_index(),
                x="UNIDAD",
                y="count",
                title="Carga de Incidencias por Unidad (Manual)",
                labels={"UNIDAD": "Unidad", "count": "Cantidad"},
                color="UNIDAD",
            )
            st.plotly_chart(fig_unidad, use_container_width=True, key=f"fig_unidad_{sufijo}")
        else:
            st.info("Asigne texto en la columna 'UNIDAD' para visualizar este gráfico.")


with tab1:
    st.subheader("📅 Comportamiento Diario por Mes")
    if df is not None:
        col_s1, col_s2, _ = st.columns([1, 1, 2])
        with col_s1:
            anio_seleccionado = st.selectbox(
                "Seleccione el Año:",
                [2025, 2026, 2027],
                index=1,
                key="anio_s1",
            )
        with col_s2:
            meses_dict = {
                1: "Enero",
                2: "Febrero",
                3: "Marzo",
                4: "Abril",
                5: "Mayo",
                6: "Junio",
                7: "Julio",
                8: "Agosto",
                9: "Septiembre",
                10: "Octubre",
                11: "Noviembre",
                12: "Diciembre",
            }
            mes_nombre_seleccionado = st.selectbox(
                "Seleccione el Mes:",
                list(meses_dict.values()),
                index=7,
                key="mes_s1",
            )
            mes_seleccionado = [
                k
                for k, v in meses_dict.items()
                if v == mes_nombre_seleccionado
            ][0]

        inicio_mes_dinamico = pd.Timestamp(
            year=anio_seleccionado, month=mes_seleccionado, day=1
        )
        ultimo_dia = calendar.monthrange(
            anio_seleccionado, mes_seleccionado
        )[1]
        fin_mes_dinamico = pd.Timestamp(
            year=anio_seleccionado,
            month=mes_seleccionado,
            day=ultimo_dia,
        ) + pd.Timedelta(days=1)

        acumulado_mes = int(
            (
                (df["RECEPCION_DT"] < inicio_mes_dinamico)
                & (
                    df["FINALIZACION_DT"].isna()
                    | (df["FINALIZACION_DT"] >= inicio_mes_dinamico)
                )
            ).sum()
        )
        recibidos_mes = int(
            (
                (df["RECEPCION_DT"] >= inicio_mes_dinamico)
                & (df["RECEPCION_DT"] < fin_mes_dinamico)
            ).sum()
        )
        total_mes = acumulado_mes + recibidos_mes
        finalizados_mes = int(
            (
                (df["FINALIZACION_DT"] >= inicio_mes_dinamico)
                & (df["FINALIZACION_DT"] < fin_mes_dinamico)
            ).sum()
        )
        pendientes_mes = total_mes - finalizados_mes

        dias_mes = pd.date_range(
            start=inicio_mes_dinamico,
            end=pd.Timestamp(
                year=anio_seleccionado,
                month=mes_seleccionado,
                day=ultimo_dia,
            ),
            freq="D",
        )
        datos_diarios = []

        for dia in dias_mes:
            inicio_dia = dia
            fin_dia = dia + pd.Timedelta(days=1)
            nombre_dia = dia.strftime("%d/%m/%Y")

            cant_recibidos = int(
                (
                    (df["RECEPCION_DT"] >= inicio_dia)
                    & (df["RECEPCION_DT"] < fin_dia)
                ).sum()
            )
            cant_acumulada = int(
                (
                    (df["RECEPCION_DT"] < inicio_dia)
                    & (
                        df["FINALIZACION_DT"].isna()
                        | (df["FINALIZACION_DT"] >= inicio_dia)
                    )
                ).sum()
            )
            total_rep = cant_recibidos + cant_acumulada
            cant_finalizados = int(
                (
                    (df["FINALIZACION_DT"] >= inicio_dia)
                    & (df["FINALIZACION_DT"] < fin_dia)
                ).sum()
            )
            efectividad_dia = (
                (cant_finalizados / total_rep * 100)
                if total_rep > 0
                else 0.0
            )

            datos_diarios.append({
                "FECHA": nombre_dia,
                "REPORTES RECIBIDOS": cant_recibidos,
                "REPORTES ACUMULADOS AL INICIAR": cant_acumulada,
                "TOTAL REPORTES": total_rep,
                "REPORTES FINALIZADOS": cant_finalizados,
                "EFECTIVIDAD (%)": round(efectividad_dia, 2),
            })

        df_dia = pd.DataFrame(datos_diarios)
        efectividad_promedio_mes = (
            df_dia["EFECTIVIDAD (%)"].mean()
            if not df_dia.empty
            else 0.0
        )

        st.markdown(
            f"##### 📌 Resumen del Mes de {mes_nombre_seleccionado}"
            f" {anio_seleccionado}"
        )
        c1, c2, c3, c4, c5, c6 = st.columns(6)
        c1.metric("📦 Acumulado Inicial", f"{acumulado_mes:,}")
        c2.metric("📥 Recibidos", f"{recibidos_mes:,}")
        c3.metric("📊 Total Incidencias", f"{total_mes:,}")
        c4.metric("✅ Finalizados", f"{finalizados_mes:,}")
        c5.metric("⏳ Pendientes", f"{pendientes_mes:,}")
        c6.metric(
            "🎯 Efectividad Prom.", f"{efectividad_promedio_mes:.1f}%"
        )
        st.divider()

        fig1 = go.Figure()
        fig1.add_trace(
            go.Bar(
                x=df_dia["FECHA"],
                y=df_dia["REPORTES ACUMULADOS AL INICIAR"],
                name="Acumulados al Iniciar",
                marker_color="#d8e4fc",
                text=df_dia["REPORTES ACUMULADOS AL INICIAR"],
                textposition="inside",
                textfont=dict(color="black", size=11),
            )
        )
        fig1.add_trace(
            go.Bar(
                x=df_dia["FECHA"],
                y=df_dia["REPORTES RECIBIDOS"],
                name="Reportes Recibidos",
                marker_color="#e9f056",
                text=df_dia["REPORTES RECIBIDOS"],
                textposition="inside",
                textfont=dict(color="black", size=11),
            )
        )
        fig1.add_trace(
            go.Scatter(
                x=df_dia["FECHA"],
                y=df_dia["TOTAL REPORTES"],
                name="Total Reportes",
                mode="text+markers",
                text=df_dia["TOTAL REPORTES"],
                textposition="top center",
                textfont=dict(color="#1F4E78", size=12),
                marker=dict(size=8, color="rgba(0,0,0,0)"),
                showlegend=False,
            )
        )
        fig1.add_trace(
            go.Scatter(
                x=df_dia["FECHA"],
                y=df_dia["REPORTES FINALIZADOS"],
                name="Reportes Finalizados",
                mode="lines+markers+text",
                text=df_dia["REPORTES FINALIZADOS"],
                textposition="bottom center",
                textfont=dict(color="#1d4ed8", size=11),
                marker=dict(size=6, color="#1d4ed8"),
                line=dict(color="#1d4ed8", width=2),
            )
        )

        fig1.update_layout(
            title=dict(
                text=(
                    "<b>Flujo Diario de Reportes -"
                    f" {mes_nombre_seleccionado} {anio_seleccionado}</b>"
                ),
                font=dict(size=18, color="#1f4e78"),
            ),
            barmode="stack",
            xaxis_title="<b>Día del Mes</b>",
            yaxis_title="<b>Cantidad de Reportes</b>",
            hovermode="x unified",
            legend=dict(
                orientation="h",
                yanchor="bottom",
                y=1.05,
                xanchor="right",
                x=1,
            ),
        )
        st.plotly_chart(fig1, use_container_width=True)

        with st.expander(
            "Ver y seleccionar incidencias del mes para análisis"
        ):
            df_mes_filtro = df[
                (df["RECEPCION_DT"] >= inicio_mes_dinamico)
                & (df["RECEPCION_DT"] < fin_mes_dinamico)
            ].copy()
            df_mes_filtro["ORIGEN_TIPO"] = "Recibidas en el Día"
            renderizar_tabla_con_seleccion(
                df_mes_filtro, "diario_mes", "Recibidas en el Día"
            )
    else:
        st.info("No hay datos cargados todavía.")


with tab2:
    st.subheader("📈 Seguimiento Mensual")
    if df is not None:
        inicio_anio = pd.Timestamp("2026-01-01")
        fin_anio = pd.Timestamp("2026-09-01")

        acumulado_anio = int(
            (
                (df["RECEPCION_DT"] < inicio_anio)
                & (
                    df["FINALIZACION_DT"].isna()
                    | (df["FINALIZACION_DT"] >= inicio_anio)
                )
            ).sum()
        )
        recibidos_anio = int(
            (
                (df["RECEPCION_DT"] >= inicio_anio)
                & (df["RECEPCION_DT"] < fin_anio)
            ).sum()
        )
        total_anio = acumulado_anio + recibidos_anio
        finalizados_anio = int(
            (
                (df["FINALIZACION_DT"] >= inicio_anio)
                & (df["FINALIZACION_DT"] < fin_anio)
            ).sum()
        )
        pendientes_anio = total_anio - finalizados_anio
        efectividad_global_anio = (
            (finalizados_anio / total_anio * 100)
            if total_anio > 0
            else 0.0
        )

        st.markdown(
            "##### 📌 Resumen Acumulado Anual (Enero - Agosto 2026)"
        )
        ac1, ac2, ac3, ac4, ac5, ac6 = st.columns(6)
        ac1.metric("📦 Acumulado Inicial", f"{acumulado_anio:,}")
        ac2.metric("📥 Recibidos", f"{recibidos_anio:,}")
        ac3.metric("📊 Total Incidencias", f"{total_anio:,}")
        ac4.metric("✅ Finalizados", f"{finalizados_anio:,}")
        ac5.metric("⏳ Pendientes", f"{pendientes_anio:,}")
        ac6.metric(
            "🎯 Efectividad Global", f"{efectividad_global_anio:.1f}%"
        )
        st.divider()

        meses_2026_hasta_agosto = [
            (1, "ENERO"),
            (2, "FEBRERO"),
            (3, "MARZO"),
            (4, "ABRIL"),
            (5, "MAYO"),
            (6, "JUNIO"),
            (7, "JULIO"),
            (8, "AGOSTO"),
        ]
        datos_anual = []

        for num_mes, nombre_mes in meses_2026_hasta_agosto:
            inicio_mes = pd.Timestamp(year=2026, month=num_mes, day=1)
            fin_mes = (
                pd.Timestamp(year=2027, month=1, day=1)
                if num_mes == 12
                else pd.Timestamp(year=2026, month=num_mes + 1, day=1)
            )

            cant_recibidos = int(
                (
                    (df["RECEPCION_DT"] >= inicio_mes)
                    & (df["RECEPCION_DT"] < fin_mes)
                ).sum()
            )
            cant_acumulada = int(
                (
                    (df["RECEPCION_DT"] < inicio_mes)
                    & (
                        df["FINALIZACION_DT"].isna()
                        | (df["FINALIZACION_DT"] >= inicio_mes)
                    )
                ).sum()
            )
            total_rep = cant_recibidos + cant_acumulada
            cant_finalizados = int(
                (
                    (df["FINALIZACION_DT"] >= inicio_mes)
                    & (df["FINALIZACION_DT"] < fin_mes)
                ).sum()
            )

            datos_anual.append({
                "MES": nombre_mes,
                "AÑO": 2026,
                "REPORTES RECIBIDOS": cant_recibidos,
                "REPORTES ACUMULADOS AL INICIAR": cant_acumulada,
                "TOTAL REPORTES": total_rep,
                "REPORTES FINALIZADOS": cant_finalizados,
            })

        df_anual = pd.DataFrame(datos_anual)

        fig2 = go.Figure()
        fig2.add_trace(
            go.Bar(
                x=df_anual["MES"],
                y=df_anual["REPORTES ACUMULADOS AL INICIAR"],
                name="Acumulados al Iniciar",
                marker_color="#d8e4fc",
                text=df_anual["REPORTES ACUMULADOS AL INICIAR"],
                textposition="inside",
                textfont=dict(color="black", size=14),
            )
        )
        fig2.add_trace(
            go.Bar(
                x=df_anual["MES"],
                y=df_anual["REPORTES RECIBIDOS"],
                name="Reportes Recibidos",
                marker_color="#e9f056",
                text=df_anual["REPORTES RECIBIDOS"],
                textposition="inside",
                textfont=dict(color="black", size=14),
            )
        )
        fig2.add_trace(
            go.Scatter(
                x=df_anual["MES"],
                y=df_anual["TOTAL REPORTES"],
                name="Total Reportes",
                mode="text+markers",
                text=df_anual["TOTAL REPORTES"],
                textposition="top center",
                textfont=dict(color="#1F4E78", size=14),
                marker=dict(size=8, color="rgba(0,0,0,0)"),
                showlegend=False,
            )
        )
        fig2.add_trace(
            go.Scatter(
                x=df_anual["MES"],
                y=df_anual["REPORTES FINALIZADOS"],
                name="Reportes Finalizados",
                mode="lines+markers+text",
                text=df_anual["REPORTES FINALIZADOS"],
                textposition="bottom center",
                textfont=dict(color="#1d4ed8", size=14),
                marker=dict(size=6, color="#1d4ed8"),
                line=dict(color="#1d4ed8", width=2),
            )
        )

        fig2.update_layout(
            title=dict(
                text=(
                    "<b>Resumen Acumulado Mensual (Enero - Agosto 2026)</b>"
                ),
                font=dict(size=18, color="#1F4E78"),
            ),
            barmode="stack",
            xaxis_title="<b>Mes</b>",
            yaxis_title="<b>Cantidad de Reportes</b>",
            hovermode="x unified",
            legend=dict(
                orientation="h",
                yanchor="bottom",
                y=1.05,
                xanchor="right",
                x=1,
            ),
        )
        st.plotly_chart(fig2, use_container_width=True)

        with st.expander(
            "Ver y seleccionar incidencias del acumulado anual"
        ):
            df_anual_filtro = df[
                (df["RECEPCION_DT"] >= inicio_anio)
                & (df["RECEPCION_DT"] < fin_anio)
            ].copy()
            df_anual_filtro["ORIGEN_TIPO"] = "Recibidas en el Día"
            renderizar_tabla_con_seleccion(
                df_anual_filtro, "anual", "Recibidas en el Día"
            )
    else:
        st.info("No hay datos cargados todavía.")


with tab3:
    st.subheader(
        "🔍 Buscador de Incidencias"
    )
    if df is not None:
        col_f1, col_f2, col_f3 = st.columns(3)
        with col_f1:
            fecha_busqueda = st.date_input(
                "Seleccione la fecha a auditar:",
                value=datetime(2026, 8, 12),
            )
        with col_f2:
            filtro_num = st.text_input(
                "Filtrar por Nº de Incidencia:", placeholder="Ej. 12345"
            )
        with col_f3:
            filtro_dir = st.text_input(
                "Filtrar por Dirección:", placeholder="Ej. Av. Principal"
            )

        inicio_sel = pd.Timestamp(fecha_busqueda)
        fin_sel = inicio_sel + pd.Timedelta(days=1)

        df_base_filtrada = df.copy()
        if filtro_num:
            df_base_filtrada = df_base_filtrada[
                df_base_filtrada["INCIDENCIA"].str.contains(
                    filtro_num, case=False, na=False
                )
            ]
        if filtro_dir:
            df_base_filtrada = df_base_filtrada[
                df_base_filtrada["DIRECCIÓN"].str.contains(
                    filtro_dir, case=False, na=False
                )
            ]

        df_acumulados_dia = df_base_filtrada[
            (df_base_filtrada["RECEPCION_DT"] < inicio_sel)
            & (
                df_base_filtrada["FINALIZACION_DT"].isna()
                | (df_base_filtrada["FINALIZACION_DT"] >= inicio_sel)
            )
        ].copy()
        df_acumulados_dia["ORIGEN_TIPO"] = "Acumuladas"

        df_recibidos_dia = df_base_filtrada[
            (df_base_filtrada["RECEPCION_DT"] >= inicio_sel)
            & (df_base_filtrada["RECEPCION_DT"] < fin_sel)
        ].copy()
        df_recibidos_dia["ORIGEN_TIPO"] = "Recibidas en el Día"

        df_finalizados_dia = df_base_filtrada[
            (df_base_filtrada["FINALIZACION_DT"] >= inicio_sel)
            & (df_base_filtrada["FINALIZACION_DT"] < fin_sel)
        ].copy()
        df_finalizados_dia["ORIGEN_TIPO"] = df_finalizados_dia[
            "RECEPCION_DT"
        ].apply(
            lambda x: "Acumuladas" if x < inicio_sel else "Recibidas en el Día"
        )

        st.markdown("---")
        st.markdown(
            f"##### 📊 Resumen para el día:"
            f" {fecha_busqueda.strftime('%d/%m/%Y')}"
        )

        total_activos_dia = len(df_acumulados_dia) + len(df_recibidos_dia)
        total_finalizados_dia = len(df_finalizados_dia)
        efectividad_dia_busqueda = (
            (total_finalizados_dia / total_activos_dia * 100)
            if total_activos_dia > 0
            else 0.0
        )
        pendientes_dia_busqueda = total_activos_dia - total_finalizados_dia

        m1, m2, m3, m4, m5, m6 = st.columns(6)
        m1.metric("📦 Acumulados Previos", f"{len(df_acumulados_dia):,}")
        m2.metric("📥 Recibidos", f"{len(df_recibidos_dia):,}")
        m3.metric("📊 Total Activos", f"{total_activos_dia:,}")
        m4.metric("✅ Finalizados", f"{total_finalizados_dia:,}")
        m5.metric("⏳ Pendientes", f"{pendientes_dia_busqueda:,}")
        m6.metric("🎯 Efectividad", f"{efectividad_dia_busqueda:.1f}%")
        st.markdown("---")

        with st.expander(
            f"📦 1. Incidencias Acumuladas Pendientes"
            f" ({len(df_acumulados_dia)} registros)",
            expanded=True,
        ):
            renderizar_tabla_con_seleccion(
                df_acumulados_dia, "busqueda_acumulados", "Acumuladas"
            )

        with st.expander(
            f"📥 2. Incidencias Recibidas ({len(df_recibidos_dia)}"
            " registros)"
        ):
            renderizar_tabla_con_seleccion(
                df_recibidos_dia, "busqueda_recibidos", "Recibidas en el Día"
            )

        with st.expander(
            f"✅ 3. Incidencias Finalizadas / Atendidas"
            f" ({len(df_finalizados_dia)} registros)"
        ):
            renderizar_tabla_con_seleccion(
                df_finalizados_dia, "busqueda_finalizados", None
            )
    else:
        st.info("No hay datos cargados todavía.")


with tab4:
    # ==========================================
    # PESTAÑA 4: ANÁLISIS
    # ==========================================
    st.header("🎯 Panel de Análisis")
    
    if df is not None:
        st.markdown("### 📝 Ingreso de Incidencias por Texto")
        texto_lista = st.text_area(
            "Escribe o pega los números de incidencias separados por comas:",
            placeholder="Ej. 12345, 67890, 11223, 44556",
            key="input_texto_lote"
        )
        
        col_btn_lote1, col_btn_lote2 = st.columns(2)
        with col_btn_lote1:
            btn_cargar = st.button("🚀 Analizar", type="primary", use_container_width=True)
        with col_btn_lote2:
            btn_limpiar = st.button("🗑️ Limpiar", use_container_width=True)

        if btn_limpiar:
            st.session_state["lote_seleccionado"] = pd.DataFrame()
            st.success("¡Lote limpiado!")
            st.rerun()
        
        if btn_cargar:
            if texto_lista.strip():
                lista_ids = [x.strip().upper() for x in texto_lista.split(",") if x.strip()]
                lista_ids = [x.replace(".0", "") for x in lista_ids]
                
                df_encontrados = df[df["INCIDENCIA"].isin(lista_ids)].copy()
                
                if not df_encontrados.empty:
                    def clasificar_origen_auto(row):
                        rec = row["RECEPCION_DT"]
                        fin = row["FINALIZACION_DT"]
                        if pd.isna(fin) or (not pd.isna(rec) and not pd.isna(fin) and rec.date() != fin.date()):
                            return "Acumuladas"
                        else:
                            return "Recibidas en el Día"

                    df_encontrados["ORIGEN_TIPO"] = df_encontrados.apply(clasificar_origen_auto, axis=1)
                    if "UNIDAD" not in df_encontrados.columns: df_encontrados["UNIDAD"] = ""
                    if "GERENCIA" not in df_encontrados.columns: df_encontrados["GERENCIA"] = ""
                    
                    st.session_state["lote_seleccionado"] = df_encontrados
                    
                    ids_encontrados_set = set(df_encontrados["INCIDENCIA"].tolist())
                    ids_faltantes = [i for i in lista_ids if i not in ids_encontrados_set]
                    
                    st.success(f"¡Lote cargado con éxito! Se encontraron {len(df_encontrados)} de {len(lista_ids)} incidencias.")
                    if ids_faltantes:
                        st.warning(f"⚠️ Las siguientes incidencias no se encontraron en la base de datos: {', '.join(ids_faltantes)}")
                    st.rerun()
                else:
                    st.error("❌ No se encontró ninguna de las incidencias ingresadas en la base de datos.")
            else:
                st.warning("⚠️ Por favor, ingresa al menos un número de incidencia.")
        
        st.divider()

    df_lote = st.session_state["lote_seleccionado"]

    if df_lote.empty:
        st.info("No hay incidencias cargadas en el lote de análisis. Ingresa una lista de números en el cuadro de texto superior o selecciónalos desde las otras pestañas.")
    else:
        st.info(f"Mostrando análisis para el lote actual de **{len(df_lote)}** incidencias.")

        if "UNIDAD" not in df_lote.columns:
            df_lote["UNIDAD"] = ""
        if "GERENCIA" not in df_lote.columns:
            df_lote["GERENCIA"] = ""

        st.subheader("✍️ Asignación Manual de Unidad y Gerencia por Incidencia")
        st.markdown("💡 *Haz doble clic en las columnas **UNIDAD** o **GERENCIA** de la tabla inferior para editar.*")

        cols_edicion = ["INCIDENCIA", "MOTIVO", "ESTADO", "UNIDAD", "GERENCIA"]
        for c in cols_edicion:
            if c not in df_lote.columns:
                df_lote[c] = ""

        df_para_editar = df_lote[cols_edicion].copy()

        df_editado_manual = st.data_editor(
            df_para_editar,
            use_container_width=True,
            key="editor_manual_unidades",
            hide_index=True,
            disabled=["INCIDENCIA", "MOTIVO", "ESTADO"],
        )

        col_save1, col_save2 = st.columns([2, 2])
        with col_save1:
            if st.button("💾 Guardar", key="btn_guardar_manual", use_container_width=True):
                for idx, row in df_editado_manual.iterrows():
                    inc_id = row["INCIDENCIA"]
                    mask = df_lote["INCIDENCIA"] == inc_id
                    df_lote.loc[mask, "UNIDAD"] = str(row["UNIDAD"]).strip().upper()
                    df_lote.loc[mask, "GERENCIA"] = str(row["GERENCIA"]).strip().upper()

                st.session_state["lote_seleccionado"] = df_lote
                st.success("¡Asignaciones guardadas y gráficos actualizados!")
                st.rerun()

        with col_save2:
            nombre_guardar_input = st.text_input("Nombre para guardar este lote:", placeholder="Ej. Análisis Operativo Semana 1", key="input_nombre_lote_guardado")
            if st.button("📥 Guardar Reporte", key="btn_almacenar_lote", type="primary", use_container_width=True):
                if nombre_guardar_input.strip():
                    timestamp_lote = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
                    nombre_llave = f"{nombre_guardar_input.strip()} ({timestamp_lote})"
                    
                    st.session_state["lotes_guardados"][nombre_llave] = df_lote.copy(deep=True)
                    st.success(f"¡Análisis guardado exitosamente con el nombre '{nombre_llave}'! Puedes consultarlo en la pestaña '📄 Reportes'.")
                else:
                    st.warning("⚠️ Ingresa un nombre válido para almacenar el análisis.")

        st.divider()

        # Renderizar gráficos para el lote activo con sufijo único
        renderizar_panel_graficos(df_lote, sufijo="activo")

        st.markdown("---")
        st.subheader("📋 Detalle Completo del Reporte")
        st.dataframe(df_lote, use_container_width=True)

with tab5:
    # ==========================================
    # PESTAÑA 5: REPORTES (INDEPENDIENTES)
    # ==========================================
    st.header("📁 Gestión de Análisis y Reportes")
    st.markdown("Aquí puedes consultar los análisis que has almacenado previamente de forma totalmente independiente. Si actualizas la base de datos general, estos datos **no cambiarán** a menos que decidas reanalizarlos.")

    lotes_dict = st.session_state["lotes_guardados"]

    if not lotes_dict:
        st.info("No hay lotes guardados todavía. Ve a la pestaña **🎯 Análisis**, configura o carga tu lista y haz clic en 'Almacenar y Congelar este Análisis'.")
    else:
        nombres_lotes_disponibles = list(lotes_dict.keys())
        lote_seleccionado_key = st.selectbox("Seleccione el Lote Guardado a Consultar:", nombres_lotes_disponibles, key="select_lote_guardado_activo")

        if lote_seleccionado_key:
            df_congelado = lotes_dict[lote_seleccionado_key]

            col_acc1, col_acc2, _ = st.columns([2, 2, 3])
            with col_acc1:
                if st.button("🔄 Reanalizar", key=f"btn_reanalizar_{lote_seleccionado_key}", use_container_width=True):
                    if df is not None:
                        lista_inc_congeladas = df_congelado["INCIDENCIA"].tolist()
                        df_fresco = df[df["INCIDENCIA"].isin(lista_inc_congeladas)].copy()
                        
                        if not df_fresco.empty:
                            def clasificar_origen_auto(row):
                                rec = row["RECEPCION_DT"]
                                fin = row["FINALIZACION_DT"]
                                if pd.isna(fin) or (not pd.isna(rec) and not pd.isna(fin) and rec.date() != fin.date()):
                                    return "Acumuladas"
                                else:
                                    return "Recibidas en el Día"
                            df_fresco["ORIGEN_TIPO"] = df_fresco.apply(clasificar_origen_auto, axis=1)
                            
                            if "UNIDAD" in df_congelado.columns and "UNIDAD" in df_fresco.columns:
                                mapeo_unidades = df_congelado.set_index("INCIDENCIA")["UNIDAD"].to_dict()
                                df_fresco["UNIDAD"] = df_fresco["INCIDENCIA"].map(mapeo_unidades).fillna("")
                            if "GERENCIA" in df_congelado.columns and "GERENCIA" in df_fresco.columns:
                                mapeo_gerencias = df_congelado.set_index("INCIDENCIA")["GERENCIA"].to_dict()
                                df_fresco["GERENCIA"] = df_fresco["INCIDENCIA"].map(mapeo_gerencias).fillna("")

                            st.session_state["lotes_guardados"][lote_seleccionado_key] = df_fresco
                            st.success("¡Lote actualizado exitosamente con los datos más recientes de la base de datos!")
                            st.rerun()
                        else:
                            st.warning("⚠️ No se encontraron las incidencias de este lote en la base de datos actual.")
                    else:
                        st.error("No hay base de datos cargada.")

            with col_acc2:
                if st.button("🗑️ Eliminar Reporte", key=f"btn_eliminar_{lote_seleccionado_key}", use_container_width=True):
                    del st.session_state["lotes_guardados"][lote_seleccionado_key]
                    st.success(f"El lote '{lote_seleccionado_key}' ha sido eliminado.")
                    st.rerun()

            st.divider()
            st.info(f"Mostrando análisis congelado para: **{lote_seleccionado_key}** (Total registros: {len(df_congelado)})")

            # Renderizar los gráficos del lote congelado con sufijo único diferente
            renderizar_panel_graficos(df_congelado, sufijo="guardado")

            st.markdown("---")
            st.subheader("📋 Detalle del Lote Guardado")
            st.dataframe(df_congelado, use_container_width=True)


with tab6:
    # ==========================================
    # PESTAÑA 6: CARGA Y ACTUALIZACIÓN
    # ==========================================
    st.subheader("📁 Carga de Data")
    st.markdown("Sube aquí tus archivos `.txt` de incidencias para actualizar el sistema.")

    if "uploader_key" not in st.session_state:
        st.session_state["uploader_key"] = 0

    archivos_cargados = st.file_uploader(
        "Selecciona archivos .txt",
        type=["txt", "TXT"],
        accept_multiple_files=True,
        key=f"uploader_{st.session_state['uploader_key']}",
    )

    if archivos_cargados:
        st.info(f"✓ {len(archivos_cargados)} archivo(s) seleccionados.")

    col_c1, col_c2 = st.columns([2, 2])
    with col_c1:
        if st.button(
            "🚀 Actualizar Base de Datos",
            key="btn_ejecutar_txt_main",
            use_container_width=True,
            type="primary"
        ):
            if archivos_cargados:
                with st.spinner("Procesando archivos y consolidando..."):
                    exito = procesar_txts_seguro(archivos_cargados)
                    if exito:
                        st.cache_data.clear()
                        st.success("¡Actualización exitosa!")
                        st.session_state["uploader_key"] += 1
                        st.rerun()
            else:
                st.warning("⚠️ Sube al menos un archivo .txt")

    if os.path.exists("resultado_actualizacion.xlsx"):
        st.markdown("---")
        st.markdown("##### 📥 Descargar Resultados")
        with open("resultado_actualizacion.xlsx", "rb") as f:
            st.download_button(
                "📥 Descargar Archivo de Actualización",
                f,
                file_name="resultado_actualizacion.xlsx",
                use_container_width=True,
            )
